"""
Generador de la presentación PPTX usando la plantilla real como molde.

Carga `PROPUESTA DE FORMATO PPT.pptx`, conserva sus elementos visuales y
solo reemplaza el contenido dinámico proveniente del JSON de Gemini.

Arquitectura:
    EXCEL -> Gemini -> JSON (secciones) -> motor de composición -> PPT

El motor usa DOS layouts de contenido definidos en la plantilla:
    * Layout "tarjetas"  (diapositiva 2): título + 3 tarjetas oscuras + 3 fotos.
    * Layout "resumen"   (diapositiva 3): kicker + título + indicadores + listas.

La portada (diapositiva 1) y el cierre (diapositiva 4) se conservan tal cual,
reemplazando únicamente los datos dinámicos de la portada.
"""

import copy
import os
from io import BytesIO

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

# ── Rutas ────────────────────────────────────────────────────────────────────
_BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE_PPTX = os.path.join(_BASE, "PROPUESTA DE FORMATO PPT.pptx")
LOGO_CONTENIDO = os.path.join(_BASE, "static", "ppt", "logo_small.jpg")

# Índices de las diapositivas de la plantilla (0-based).
IDX_PORTADA = 0
IDX_TARJETAS = 1
IDX_RESUMEN = 2
IDX_CIERRE = 3

_MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
          "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]

# Posiciones (en pulgadas) medidas de la plantilla.
PORTADA_SUPERVISOR = (0.700, 2.050)
PORTADA_TITULO = (0.650, 2.420)
PORTADA_FECHA = (0.700, 3.280)
PORTADA_ETIQUETA = (1.000, 4.450)
PORTADA_TEMAS = (1.000, 4.850)

TARJETAS_TITULO = (0.650, 0.340)
TARJETAS_FOOTER = (0.650, 7.130)
TARJETAS_PAGINA = (12.250, 7.100)
TARJETAS_LOGO = (11.250, 0.220, 1.450, 0.350)
TARJETAS_CARD_LEFT = [0.450, 4.690, 8.930]
TARJETAS_CARD_TOP = 1.023
TARJETAS_TEXT_DY = [0.147, 0.417, 0.867, 1.127]  # offset y de label/titulo/fecha/desc desde card_top
TARJETAS_PHOTO_LEFT = [0.656, 4.847, 9.443]
TARJETAS_PHOTO_TOP = 3.840
TARJETAS_PHOTO_W = [3.538, 3.637, 2.923]
TARJETAS_PHOTO_H = 2.180

RESUMEN_KICKER = (0.700, 0.320)
RESUMEN_TITULO = (0.633, 0.550)
RESUMEN_LOGO = (11.250, 0.220, 1.450, 0.350)
RESUMEN_KPI_LABEL = [(0.860, 1.200), (1.690, 1.200), (2.550, 1.169)]
RESUMEN_KPI_VALUE = [(0.860, 1.338), (1.810, 1.338), (2.670, 1.338)]
RESUMEN_PROD_IMPUTADO = (8.803, 1.371)
RESUMEN_PROD_RECUPERADO = (8.803, 1.586)
RESUMEN_PROD_COL1 = (6.223, 1.346)
RESUMEN_PROD_COL2 = (7.303, 1.346)
RESUMEN_POR_TIPO_TITULO = (0.980, 2.400)
RESUMEN_POR_TALLER_TITULO = (0.780, 4.515)
RESUMEN_POR_TALLER_HEAD = (0.780, 4.790)
RESUMEN_POR_TIPO_LABEL = [(0.900, 2.820), (0.900, 3.250), (0.900, 3.680), (0.900, 4.110)]
RESUMEN_POR_TIPO_VALUE = [(2.930, 2.820), (2.930, 3.250), (2.930, 3.680), (2.930, 4.110)]
RESUMEN_POR_TALLER_LABEL = [(0.790, 4.955), (0.790, 5.142), (0.790, 5.325), (0.790, 5.518), (0.775, 5.713), (0.765, 5.894)]
RESUMEN_POR_TALLER_VALUE = [(2.930, 4.975), (2.930, 5.151), (2.965, 5.319), (2.968, 5.530), (2.965, 5.713), (2.965, 5.891)]
RESUMEN_ACT_TITULO = (4.420, 2.382)
RESUMEN_PEND_TITULO = (4.420, 5.202)
RESUMEN_PROD_TITULO = (6.223, 1.166)
RESUMEN_PROD_PERDIDA_TITULO = (8.783, 1.166)
RESUMEN_POR_TALLER_CANT = (3.055, 4.790)


def _as_list(v):
    return v if isinstance(v, list) else []


def _as_dict(v):
    return v if isinstance(v, dict) else {}


def _fmt_fecha(fecha):
    if not fecha:
        return ""
    try:
        if not hasattr(fecha, "day"):
            return str(fecha)
        return f"{fecha.day} de {_MESES[fecha.month - 1]} {fecha.year}"
    except Exception:
        return str(fecha)


def _fmt_mes_ano(fecha):
    if not fecha:
        return ""
    try:
        if not hasattr(fecha, "month"):
            return str(fecha)
        return f"{_MESES[fecha.month - 1]} {fecha.year}"
    except Exception:
        return str(fecha)


# ── Utilidades de bajo nivel sobre la plantilla cargada ─────────────────────
def _set_text(shape, text):
    """Reemplaza el texto de una forma conservando el formato del primer run."""
    if text is None:
        text = ""
    text = str(text)
    tf = shape.text_frame
    while len(tf.paragraphs) > 1:
        p = tf.paragraphs[-1]._p
        p.getparent().remove(p)
    para = tf.paragraphs[0]
    runs = list(para.runs)
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        r = para.add_run()
        r.text = text


def _clear_text(slide):
    """Blanquea el texto de todas las formas (conserva formas, posiciones y formato).

    Se usa al duplicar una diapositiva de la plantilla para eliminar TODO el
    contenido de ejemplo antes de insertar únicamente los datos del JSON."""
    for sh in slide.shapes:
        if sh.has_text_frame:
            try:
                _set_text(sh, "")
            except Exception:
                pass


def _set_text_at(slide, left, top, text, tol=0.06):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        try:
            l = Emu(sh.left).inches
            t = Emu(sh.top).inches
        except Exception:
            continue
        if abs(l - left) <= tol and abs(t - top) <= tol:
            _set_text(sh, text)
            return True
    return False


def _shapes_at(slide, left, top, tol=0.2):
    result = []
    for sh in slide.shapes:
        try:
            l = Emu(sh.left).inches
            t = Emu(sh.top).inches
        except Exception:
            continue
        if abs(l - left) <= tol and abs(t - top) <= tol:
            result.append(sh)
    return result


def _delete_slides(prs, slides_to_delete):
    """Elimina diapositivas de la presentación (quita su sldId y su relación).

    Debe usarse SOLO después de haber añadido todas las diapositivas nuevas, para
    no interferir con la asignación de partnames de python-pptx."""
    partnames = {s.part.partname for s in slides_to_delete}
    sldIdLst = prs.slides._sldIdLst
    for el in list(sldIdLst):
        rId = el.get(qn("r:id"))
        rel = prs.part.rels.get(rId)
        if rel is not None and rel.target_part.partname in partnames:
            sldIdLst.remove(el)
            prs.part.drop_rel(rId)


def _reorder_slides(prs, ordered_slides):
    """Reordena las diapositivas al orden indicado (por referencia de objeto)."""
    sldIdLst = prs.slides._sldIdLst
    rId_to_partname = {}
    for rId, rel in prs.part.rels.items():
        if rel.reltype.endswith("/slide"):
            rId_to_partname[rId] = rel.target_part.partname
    partname_to_sldId = {}
    for el in list(sldIdLst):
        rId = el.get(qn("r:id"))
        partname_to_sldId[rId_to_partname.get(rId)] = el
    for el in list(sldIdLst):
        sldIdLst.remove(el)
    for slide in ordered_slides:
        el = partname_to_sldId.get(slide.part.partname)
        if el is not None:
            sldIdLst.append(el)


def _duplicate_slide(prs, layout, shape_elements):
    """Crea una diapositiva nueva duplicando los elementos (sin imágenes) dados."""
    dest = prs.slides.add_slide(layout)
    for shp in list(dest.shapes):
        shp._element.getparent().remove(shp._element)
    for el in shape_elements:
        dest.shapes._spTree.append(copy.deepcopy(el))
    return dest


def _add_logo(slide, pos):
    x, y, w, h = pos
    if os.path.exists(LOGO_CONTENIDO):
        slide.shapes.add_picture(LOGO_CONTENIDO, Inches(x), Inches(y), Inches(w), Inches(h))


def _delete_shapes_in_region(slide, x1, y1, x2, y2):
    """Elimina las formas cuyo vértice superior-izquierdo cae dentro de la región."""
    for shp in list(slide.shapes):
        try:
            l = Emu(shp.left).inches
            t = Emu(shp.top).inches
        except Exception:
            continue
        if x1 <= l <= x2 and y1 <= t <= y2:
            shp._element.getparent().remove(shp._element)


def _add_picture_fitted(slide, data_bytes, x, y, max_w, max_h):
    if not data_bytes:
        return None
    try:
        from PIL import Image as PILImage
        with PILImage.open(BytesIO(data_bytes)) as im:
            iw, ih = im.size
    except Exception:
        iw, ih = 4, 3
    if not iw or not ih:
        iw, ih = 4, 3
    ratio = iw / ih
    w = max_w
    h = w / ratio
    if h > max_h:
        h = max_h
        w = h * ratio
    px = x + (max_w - w) / 2
    py = y + (max_h - h) / 2
    return slide.shapes.add_picture(BytesIO(data_bytes), Inches(px), Inches(py), Inches(w), Inches(h))


def _imagen_por_indice(imagenes_archivo, indice):
    if indice is None:
        return None
    try:
        i = int(indice) - 1
    except (TypeError, ValueError):
        return None
    if 0 <= i < len(imagenes_archivo):
        return imagenes_archivo[i].get("bytes")
    return None


# ── Relleno de la portada ────────────────────────────────────────────────────
def _fill_portada(slide, informe, data):
    portada = _as_dict(data.get("portada"))
    reg = _as_dict(data.get("registro_guardia"))

    supervisor = (informe.get("fiscalizador") or "").strip().upper()
    _set_text_at(slide, *PORTADA_SUPERVISOR,
                 f"SUPERVISOR. {supervisor}" if supervisor else "SUPERVISOR.")

    titulo = portada.get("titulo") or "Reporte de Guardia"
    _set_text_at(slide, *PORTADA_TITULO, titulo)

    fecha = _fmt_fecha(informe.get("fecha_guardia"))
    _set_text_at(slide, *PORTADA_FECHA, fecha)

    etiqueta = portada.get("etiqueta") or "Talleres supervisados"
    _set_text_at(slide, *PORTADA_ETIQUETA, etiqueta)

    temas = _as_list(portada.get("temas"))
    if not temas:
        nombres = []
        for t in _as_list(reg.get("talleres")):
            n = (_as_dict(t).get("taller") or "").strip()
            if n and n not in nombres:
                nombres.append(n)
        for p in _as_list(reg.get("pendientes")):
            n = (_as_dict(p).get("taller") or "").strip()
            if n and n not in nombres:
                nombres.append(n)
        temas = nombres
    if temas:
        _set_text_at(slide, *PORTADA_TEMAS, " · ".join(temas))


# ── Layout "tarjetas" (evidencias / incidentes) ─────────────────────────────
def _evidencia_titulo(ev):
    ev = _as_dict(ev)
    t = (ev.get("titulo") or "").strip()
    if not t:
        t = (ev.get("pozo_equipo") or "").strip()
    if not t:
        t = (ev.get("descripcion") or "").strip()[:60]
    return t


def _evidencia_fecha(ev):
    ev = _as_dict(ev)
    return ev.get("fecha_hora") or ev.get("fecha") or ""


def _evidencia_descripcion(ev):
    ev = _as_dict(ev)
    partes = []
    if ev.get("descripcion"):
        partes.append(str(ev["descripcion"]))
    if ev.get("hallazgos"):
        partes.append(str(ev["hallazgos"]))
    return "\n".join(partes)


def _fill_tarjetas(slide, titulo, evidencias, fotos, mes_ano, num):
    _clear_text(slide)
    _set_text_at(slide, *TARJETAS_TITULO, titulo)
    _set_text_at(slide, *TARJETAS_FOOTER, f"Reporte de Guardia | {mes_ano}".strip(" |"))
    _set_text_at(slide, *TARJETAS_PAGINA, str(num).zfill(2))
    _add_logo(slide, TARJETAS_LOGO)

    for i in range(3):
        card_left = TARJETAS_CARD_LEFT[i]
        ev = evidencias[i] if i < len(evidencias) else None
        foto = fotos[i] if i < len(fotos) else None

        if ev is None:
            # Eliminar la tarjeta (y su zona de foto) para no dejar cajas oscuras vacías.
            _delete_shapes_in_region(slide, card_left - 0.05, TARJETAS_CARD_TOP - 0.05,
                                     card_left + 3.95, TARJETAS_CARD_TOP + 2.55)
            _delete_shapes_in_region(slide, TARJETAS_PHOTO_LEFT[i] - 0.05, TARJETAS_PHOTO_TOP - 0.05,
                                     TARJETAS_PHOTO_LEFT[i] + TARJETAS_PHOTO_W[i] + 0.05,
                                     TARJETAS_PHOTO_TOP + TARJETAS_PHOTO_H + 0.05)
            continue

        # Rellenar los 4 textos de la tarjeta.
        label_top = TARJETAS_CARD_TOP + TARJETAS_TEXT_DY[0]
        titulo_top = TARJETAS_CARD_TOP + TARJETAS_TEXT_DY[1]
        fecha_top = TARJETAS_CARD_TOP + TARJETAS_TEXT_DY[2]
        desc_top = TARJETAS_CARD_TOP + TARJETAS_TEXT_DY[3]

        ev = _as_dict(ev)
        _set_text_at(slide, card_left + 0.200, label_top,
                     (ev.get("tipo") or ev.get("clasificacion") or "").upper())
        _set_text_at(slide, card_left + 0.200, titulo_top, _evidencia_titulo(ev))
        _set_text_at(slide, card_left + 0.200, fecha_top, _evidencia_fecha(ev))
        _set_text_at(slide, card_left + 0.200, desc_top, _evidencia_descripcion(ev))

        if foto:
            _add_picture_fitted(slide, foto,
                                TARJETAS_PHOTO_LEFT[i], TARJETAS_PHOTO_TOP,
                                TARJETAS_PHOTO_W[i], TARJETAS_PHOTO_H)


def _render_evidencias(prs, tarjetas_layout, tarjetas_shapes, mes_ano, num, titulo, evidencias, imagenes_archivo):
    evidencias = _as_list(evidencias)
    if not evidencias:
        return num
    por_slide = 3
    for i in range(0, len(evidencias), por_slide):
        chunk = evidencias[i:i + por_slide]
        fotos = []
        for ev in chunk:
            fotos.append(_imagen_por_indice(imagenes_archivo, _as_dict(ev).get("imagen_indice")))
        slide = _duplicate_slide(prs, tarjetas_layout, tarjetas_shapes)
        _fill_tarjetas(slide, titulo, chunk, fotos, mes_ano, num)
        num += 1
    return num


# ── Layout "resumen" (indicadores y listas) ─────────────────────────────────
def _add_items(slide, x, y, items, size=11, color=RGBColor(0x10, 0x2A, 0x43)):
    """Añade ítems de lista con el estilo de la plantilla dentro de una tarjeta."""
    items = _as_list(items)
    for it in items:
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(7.8), Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = "• " + str(it)
        r.font.size = Pt(size)
        r.font.name = "Aptos"
        r.font.color.rgb = color
        y += 0.32


def _fill_resumen(slide, titulo, seccion, data, mes_ano, num):
    _clear_text(slide)
    _set_text_at(slide, *RESUMEN_KICKER, "")
    _set_text_at(slide, *RESUMEN_TITULO, titulo)
    _set_text_at(slide, *TARJETAS_FOOTER, f"Reporte de Guardia | {mes_ano}".strip(" |"))
    _set_text_at(slide, *TARJETAS_PAGINA, str(num).zfill(2))
    _add_logo(slide, RESUMEN_LOGO)

    # Indicadores (máx. 3)
    indicadores = _as_list(seccion.get("indicadores"))
    for i in range(3):
        ind = _as_dict(indicadores[i]) if i < len(indicadores) else {}
        nombre = ind.get("nombre") or ""
        valor = ind.get("valor") if ind.get("valor") not in (None, "") else ""
        _set_text_at(slide, *RESUMEN_KPI_LABEL[i], nombre.upper())
        _set_text_at(slide, *RESUMEN_KPI_VALUE[i], str(valor))

    # Producción (PRODUCCIÓN + PROD. PERDIDA). Se oculta si no hay datos.
    produccion = _as_dict(seccion.get("produccion"))
    if not any(produccion.values()):
        produccion = _as_dict(data.get("produccion"))
    if not any(produccion.values()):
        produccion = _as_dict(_as_dict(data.get("registro_guardia")).get("cabecera"))
    oil = produccion.get("produccion_oil_bpd")
    gas = produccion.get("produccion_gas_mpc")
    imputado = produccion.get("prod_perdida_imputada")
    recuperado = produccion.get("produccion_recuperada")

    hay_produccion = any(v not in (None, "") for v in (oil, gas, imputado, recuperado))
    if hay_produccion:
        _set_text_at(slide, *RESUMEN_PROD_TITULO, "PRODUCCIÓN")
        _set_text_at(slide, *RESUMEN_PROD_PERDIDA_TITULO, "PROD. PERDIDA")
        _set_text_at(slide, *RESUMEN_PROD_COL1, f"Bls\n{oil if oil not in (None, '') else '—'}")
        _set_text_at(slide, *RESUMEN_PROD_COL2, f"Mpc\n{gas if gas not in (None, '') else '—'}")
        _set_text_at(slide, *RESUMEN_PROD_IMPUTADO,
                     f"Imputado: {imputado}" if imputado not in (None, "") else "Imputado: —")
        _set_text_at(slide, *RESUMEN_PROD_RECUPERADO,
                     f"Recuperado: {recuperado}" if recuperado not in (None, "") else "Recuperado: —")
    else:
        # Ocultar las tarjetas PRODUCCIÓN y PROD. PERDIDA (sin datos de producción).
        _delete_shapes_in_region(slide, 6.0, 1.0, 10.9, 2.0)

    # Distribución por tipo
    tipos = _as_list(seccion.get("tipos"))
    _set_text_at(slide, *RESUMEN_POR_TIPO_TITULO, "POR TIPO")
    # Eliminar las barras de ejemplo de "POR TIPO" (no se actualizan dinámicamente).
    _delete_shapes_in_region(slide, 0.85, 2.98, 2.60, 4.35)
    for i, pos in enumerate(RESUMEN_POR_TIPO_LABEL):
        item = _as_dict(tipos[i]) if i < len(tipos) else {}
        nombre = item.get("nombre") or ""
        valor = item.get("valor") or ""
        pct = item.get("porcentaje") or ""
        if pct:
            valor = f"{valor} ({pct})"
        _set_text_at(slide, *pos, nombre)
        _set_text_at(slide, *RESUMEN_POR_TIPO_VALUE[i], valor)

    # Distribución por taller
    talleres = _as_list(seccion.get("talleres"))
    _set_text_at(slide, *RESUMEN_POR_TALLER_TITULO, "POR TALLER")
    _set_text_at(slide, *RESUMEN_POR_TALLER_HEAD, "TALLER")
    _set_text_at(slide, *RESUMEN_POR_TALLER_CANT, "CANT.")
    for i, pos in enumerate(RESUMEN_POR_TALLER_LABEL):
        item = _as_dict(talleres[i]) if i < len(talleres) else {}
        nombre = item.get("nombre") or ""
        valor = item.get("valor") or ""
        _set_text_at(slide, *pos, nombre)
        _set_text_at(slide, *RESUMEN_POR_TALLER_VALUE[i], valor)

    # Actividades relevantes
    actividades = _as_list(seccion.get("actividades"))
    if actividades:
        _set_text_at(slide, *RESUMEN_ACT_TITULO, "Actividades relevantes:")
        _add_items(slide, 4.60, 2.80, actividades)

    # Pendientes relevantes
    pendientes = _as_list(seccion.get("pendientes"))
    if pendientes:
        _set_text_at(slide, *RESUMEN_PEND_TITULO, "Pendientes (relevantes):")
        _add_items(slide, 4.60, 5.60, pendientes)


def _render_resumen(prs, resumen_layout, resumen_shapes, mes_ano, num, seccion, data):
    slide = _duplicate_slide(prs, resumen_layout, resumen_shapes)
    titulo = seccion.get("titulo") or "Resumen ejecutivo"
    _fill_resumen(slide, titulo, seccion, data, mes_ano, num)
    return num + 1


# ── Derivar secciones desde registro_guardia (respaldo de caché viejo) ──────
def _secciones_desde_registro(data):
    reg = _as_dict(data.get("registro_guardia"))
    secciones = []

    cabecera = _as_dict(reg.get("cabecera"))
    talleres = _as_list(reg.get("talleres"))
    evidencias = _as_list(reg.get("evidencias"))
    pendientes_reg = _as_list(reg.get("pendientes"))

    indicadores = []
    for lbl, key in (("Producción OIL (BPD)", "produccion_oil_bpd"),
                     ("Prod. Perdida imputada", "prod_perdida_imputada"),
                     ("Producción GAS (MPC)", "produccion_gas_mpc"),
                     ("Producción recuperada", "produccion_recuperada")):
        v = cabecera.get(key)
        if v not in (None, ""):
            indicadores.append({"nombre": lbl, "valor": v})

    if indicadores or talleres or pendientes_reg:
        secciones.append({
            "titulo": "Resumen ejecutivo",
            "tipo": "resumen",
            "indicadores": indicadores[:3],
            "talleres": [{"nombre": _as_dict(t).get("taller") or "", "valor": ""} for t in talleres],
            "actividades": [],
            "pendientes": [],
        })

    if evidencias:
        secciones.append({
            "titulo": "Seguridad y Medio Ambiente",
            "tipo": "evidencias",
            "evidencias": evidencias,
        })

    return secciones


# ── Generador principal ─────────────────────────────────────────────────────
def generar_informe_ppt(informe, data, imagenes_archivo=None):
    """Genera la presentación a partir de la plantilla real y el contenido de Gemini."""
    imagenes_archivo = imagenes_archivo or []
    data = data or {}

    if not os.path.exists(TEMPLATE_PPTX):
        raise FileNotFoundError(f"Plantilla PPT no encontrada: {TEMPLATE_PPTX}")

    prs = Presentation(TEMPLATE_PPTX)
    slides = list(prs.slides)
    if len(slides) < 4:
        raise ValueError("La plantilla debe tener al menos 4 diapositivas (portada, contenido x2, cierre).")

    portada = slides[IDX_PORTADA]
    tarjetas_src = slides[IDX_TARJETAS]
    resumen_src = slides[IDX_RESUMEN]
    cierre = slides[IDX_CIERRE]

    # Capturar el layout y los elementos (sin imágenes) de las 2 diapositivas de contenido.
    tarjetas_layout = tarjetas_src.slide_layout
    tarjetas_shapes = [copy.deepcopy(sh._element) for sh in tarjetas_src.shapes
                       if sh.shape_type != MSO_SHAPE_TYPE.PICTURE]
    resumen_layout = resumen_src.slide_layout
    resumen_shapes = [copy.deepcopy(sh._element) for sh in resumen_src.shapes
                      if sh.shape_type != MSO_SHAPE_TYPE.PICTURE]

    _fill_portada(portada, informe, data)

    secciones = _as_list(data.get("secciones"))
    if not secciones:
        secciones = _secciones_desde_registro(data)

    mes_ano = _fmt_mes_ano(informe.get("fecha_guardia"))

    # Construir todas las diapositivas de contenido (siempre duplicando la plantilla,
    # nunca reutilizando en sitio), en el orden de las secciones.
    num = 2
    contenido = []
    for sec in secciones:
        sec = _as_dict(sec)
        tipo = (sec.get("tipo") or "resumen").strip().lower()
        if tipo == "evidencias":
            evidencias = _as_list(sec.get("evidencias"))
            titulo = sec.get("titulo") or "Seguridad y Medio Ambiente"
            for i in range(0, len(evidencias), 3):
                chunk = evidencias[i:i + 3]
                slide = _duplicate_slide(prs, tarjetas_layout, tarjetas_shapes)
                fotos = [_imagen_por_indice(imagenes_archivo, _as_dict(ev).get("imagen_indice")) for ev in chunk]
                _fill_tarjetas(slide, titulo, chunk, fotos, mes_ano, num)
                num += 1
                contenido.append(slide)
        else:
            slide = _duplicate_slide(prs, resumen_layout, resumen_shapes)
            titulo = sec.get("titulo") or "Resumen ejecutivo"
            _fill_resumen(slide, titulo, sec, data, mes_ano, num)
            num += 1
            contenido.append(slide)

    # Eliminar las dos diapositivas de contenido originales de la plantilla.
    _delete_slides(prs, [tarjetas_src, resumen_src])

    # Reordenar: portada -> contenido -> cierre.
    _reorder_slides(prs, [portada] + contenido + [cierre])

    return prs
